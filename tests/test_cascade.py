import copy
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from cascade_judge import (CascadeJudge, JudgeError, JEV, DEEPSEEK, PROVIDERS,
                           exclusive_run, validate_jev, validate_deepseek, read)
from run_evaluation import PROMPT_TEMPLATE
from aggregate_scores import compute_dimension_averages
from run_cascade_evaluation import main, collect_reports


def task(n=3):
    return {'task': 'Synthetic task', 'rubric': {'info_recall': [f'item-{i}' for i in range(n)]}}


class Fake:
    def __init__(self, choices=None, confidences=None, overrides=None):
        self.choices = choices or {}
        self.confidences = confidences or {}
        self.overrides = list(overrides or [])
        self.calls = []
        self.catalogs = 0

    def catalog(self, model):
        self.catalogs += 1
        tags = ['typesafe'] if model == JEV else PROVIDERS
        return {'data': {'endpoints': [{'tag': tag, 'pricing': {'prompt': '0', 'completion': '0'}} for tag in tags]}}

    def post(self, endpoint, payload):
        self.calls.append((endpoint, copy.deepcopy(payload)))
        if self.overrides:
            override = self.overrides.pop(0)
            if isinstance(override, Exception):
                raise override
            if override is not None:
                return override
        if payload['model'] == JEV:
            answers = {}
            for key, q in payload['questions'].items():
                item = q['instructions'].split(': ', 1)[1]
                choice = self.choices.get(item, 'met')
                answers[key] = dict(type='choice', choice=choice,
                                    probabilities={c: int(c == choice) for c in ['met', 'not_met', 'blocked']})
                if self.confidences.get(item, .9) is not None:
                    answers[key]['confidence'] = self.confidences.get(item, .9)
            raw = dict(model=JEV, provider='TypeSafe', answers=answers)
        else:
            prompt = payload['messages'][0]['content']
            rubric = json.loads(prompt.split('<task_and_rubric>')[1].split('</task_and_rubric>')[0])
            results = [dict(rubric_item=text, score=0, reason='Not present', evidence='') for text in rubric['rubric_items']]
            raw = dict(model=DEEPSEEK, provider='Nebius', choices=[dict(finish_reason='stop', message=dict(content=json.dumps({'results': results})))])
        raw['usage'] = dict(cost=.000001, total_tokens=10)
        return 200, raw


class Tests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)

    def judge(self, fake, budget=2):
        return CascadeJudge(self.root/'cache', budget, fake, PROMPT_TEMPLATE)

    def test_threshold_missing_confidence_and_aggregation(self):
        fake = Fake(confidences={'item-0': .8, 'item-1': .799, 'item-2': None})
        result = self.judge(fake).evaluate('m', 1, 'report', task())
        rows = result['scores']['info_recall']
        self.assertEqual([r['route'] for r in rows.values()], ['jev', 'deepseek', 'deepseek'])
        self.assertAlmostEqual(compute_dimension_averages(result)['total'], 1/3)
        policy = fake.calls[-1][1]['provider']
        self.assertEqual(policy, dict(only=PROVIDERS, order=PROVIDERS, allow_fallbacks=True, data_collection='deny'))

    def test_blocked_in_last_batch_reviews_whole_report(self):
        fake = Fake(choices={'item-50': 'blocked'})
        result = self.judge(fake).evaluate('m', 1, 'report', task(51))
        self.assertEqual([p['model'] for _, p in fake.calls], [JEV, JEV, DEEPSEEK, DEEPSEEK])
        self.assertTrue(all(r['route'] == 'sourceguard' for r in result['scores']['info_recall'].values()))
        self.assertTrue(all('SOURCE ATTRIBUTION CHECK:' in p['messages'][0]['content'] for _, p in fake.calls[2:]))

    def test_repack_low_confidence_across_jev_batches(self):
        fake = Fake(confidences={'item-0': .1, 'item-50': .1})
        self.judge(fake).evaluate('m', 1, 'report', task(51))
        self.assertEqual(len(fake.calls), 3)
        prompt = fake.calls[-1][1]['messages'][0]['content']
        self.assertIn('"item-0",\n    "item-50"', prompt)

    def test_resume_no_calls_and_no_double_charge(self):
        fake = Fake()
        judge = self.judge(fake)
        first = judge.evaluate('m', 1, 'report', task())
        spending = judge.spending()
        second = self.judge(fake).evaluate('m', 1, 'report', task())
        self.assertEqual(first, second)
        self.assertEqual(len(fake.calls), 1)
        self.assertEqual(judge.spending(), spending)

    def test_changed_inputs_refused(self):
        fake = Fake()
        judge = self.judge(fake)
        judge.evaluate('m', 1, 'report', task())
        with self.assertRaisesRegex(JudgeError, 'changed'):
            judge.evaluate('m', 1, 'changed report', task())
        self.assertEqual(len(fake.calls), 1)

    def test_unknown_timeout_never_retried(self):
        fake = Fake(overrides=[TimeoutError('secret must not be logged')])
        for _ in range(2):
            with self.assertRaises(JudgeError):
                self.judge(fake).evaluate('m', 1, 'report', task())
        self.assertEqual(len(fake.calls), 1)
        self.assertEqual(self.judge(fake).spending()['unknown_attempts'], 1)
        self.assertNotIn('secret', next((self.root/'cache').glob('reports/*/*attempt*')).read_text())

    def test_http200_terminal504_bounded_retry_and_unknown_bill(self):
        fake = Fake(overrides=[(200, {'error': {'code': 504}})])
        judge = self.judge(fake)
        judge.evaluate('m', 1, 'report', task())
        self.assertEqual(len(fake.calls), 2)
        self.assertEqual(judge.spending()['unknown_attempts'], 1)
        self.assertGreater(judge.spending()['upper_bound_usd'], judge.spending()['confirmed_usd'])

    def test_invalid_twice_never_resets_retry_limit(self):
        bad = (200, {'usage': {'cost': .000001}})
        fake = Fake(overrides=[bad, bad])
        for _ in range(2):
            with self.assertRaisesRegex(JudgeError, 'exhausted'):
                self.judge(fake).evaluate('m', 1, 'report', task())
        self.assertEqual(len(fake.calls), 2)
        self.assertAlmostEqual(self.judge(fake).spending()['confirmed_usd'], .000002)

    def test_budget_prevents_dispatch(self):
        fake = Fake()
        with self.assertRaisesRegex(JudgeError, 'Budget exhausted'):
            self.judge(fake, .00000001).evaluate('m', 1, 'report', task())
        self.assertFalse(fake.calls)

    def test_invalid_confidence_and_probabilities_rejected(self):
        batch = [{'text':'item-0', 'dimension':'info_recall'}]
        raw = dict(model=JEV, provider='TypeSafe', answers={'criterion_000': dict(type='choice', choice='met', probabilities={'met':1,'not_met':0,'blocked':0}, confidence=float('nan'))})
        with self.assertRaises(JudgeError):
            validate_jev(batch, raw)
        raw['answers']['criterion_000']['confidence'] = .9
        raw['answers']['criterion_000']['probabilities']['met'] = .5
        with self.assertRaises(JudgeError):
            validate_jev(batch, raw)

    def test_deepseek_rejects_duplicate_boolean_and_missing_evidence(self):
        batch = [{'text':x, 'dimension':'analysis'} for x in ['a','b']]
        for results in [
            [dict(rubric_item='a',score=0,reason='',evidence='')]*2,
            [dict(rubric_item=x,score=True,reason='',evidence='ok') for x in ['a','b']],
            [dict(rubric_item=x,score=1,reason='',evidence='') for x in ['a','b']],
        ]:
            raw = dict(model=DEEPSEEK,provider='InferenceNet',choices=[dict(finish_reason='stop',message=dict(content=json.dumps(dict(results=results))))])
            with self.assertRaises(JudgeError):
                validate_deepseek(batch,raw)

    def test_exclusive_cache_lock(self):
        with exclusive_run(self.root):
            with self.assertRaises(JudgeError):
                with exclusive_run(self.root):
                    self.fail('second writer entered')
        self.assertFalse((self.root/'run.lock').exists())

    def test_cli_resume_and_failure_does_not_overwrite_export(self):
        reports = self.root/'reports'/'m'
        reports.mkdir(parents=True)
        (reports/'idx-1.md').write_text('report')
        tasks = self.root/'tasks.jsonl'
        tasks.write_text(json.dumps(dict(idx=1, content=task()))+'\n')
        output = self.root/'out.jsonl'
        args = ['--pdf_dir',str(reports.parent),'--tasks_jsonl',str(tasks), '--out_jsonl',str(output),'--cache_dir',str(self.root/'cache')]
        fake = Fake()
        with patch('run_cascade_evaluation.OpenRouter',return_value=fake):
            self.assertEqual(main(args),0)
            self.assertEqual(main(args),0)
            self.assertEqual(len(output.read_text().splitlines()),1)
            original = output.read_bytes()
            (reports/'idx-1.md').write_text('changed')
            self.assertEqual(main(args),1)
            self.assertEqual(output.read_bytes(),original)
        self.assertEqual(len(fake.calls),1)

    def test_missing_cost_stops_and_resume_does_not_resend(self):
        fake = Fake(overrides=[(200, {'model': JEV})])
        for _ in range(2):
            with self.assertRaises(JudgeError):
                self.judge(fake).evaluate('m', 1, 'report', task())
        self.assertEqual(len(fake.calls), 1)
        self.assertEqual(self.judge(fake).spending()['unknown_attempts'], 1)

    def test_auth_error_is_not_retried(self):
        fake = Fake(overrides=[(401, {'error': {'code': 401}})])
        for _ in range(2):
            with self.assertRaises(JudgeError):
                self.judge(fake).evaluate('m', 1, 'report', task())
        self.assertEqual(len(fake.calls), 1)

    def test_price_increase_prevents_dispatch(self):
        fake = Fake()
        fake.catalog = lambda model: {'data': {'endpoints': [
            {'tag':'typesafe','pricing':{'prompt':'1','completion':'0'}}]}}
        with self.assertRaisesRegex(JudgeError, 'price exceeds'):
            self.judge(fake).evaluate('m',1,'report',task())
        self.assertFalse(fake.calls)

    def test_configuration_change_cannot_mix_judges(self):
        self.judge(Fake())
        with self.assertRaisesRegex(JudgeError, 'configuration changed'):
            CascadeJudge(self.root/'cache',2,Fake(),PROMPT_TEMPLATE+'changed')

    def test_docx_input(self):
        from docx import Document
        reports = self.root/'reports'/'m'; reports.mkdir(parents=True)
        doc = Document(); doc.add_paragraph('Synthetic report')
        table = doc.add_table(rows=2, cols=1)
        table.cell(0,0).text='Column'; table.cell(1,0).text='Value'
        doc.save(reports/'idx-1.docx')
        tasks = self.root/'tasks.jsonl'; tasks.write_text(json.dumps(dict(idx=1,content=task())))
        collected = collect_reports(reports.parent,tasks)
        self.assertIn('Synthetic report',collected[0][2])
        self.assertIn('| Value |',collected[0][2])

    def test_duplicate_report_files_rejected(self):
        reports = self.root/'reports'/'m'; reports.mkdir(parents=True)
        (reports/'idx-1.md').write_text('report')

        tasks = self.root/'tasks.jsonl'; tasks.write_text(json.dumps(dict(idx=1,content=task())))
        from docx import Document
        doc = Document(); doc.add_paragraph('report'); doc.save(reports/'idx-1.docx')
        with self.assertRaises(JudgeError):
            collect_reports(reports.parent,tasks)


if __name__ == '__main__':
    unittest.main()
